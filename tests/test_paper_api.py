"""Paper trading API tests."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

import app as app_module
from core.paper_report import PaperReportService
from core.paper_trading import PaperTradingService
from tests.test_paper_trading import _config, _quote, _service


def _patch_paper_services(monkeypatch, service: PaperTradingService) -> None:
    monkeypatch.setattr(app_module, "paper_trading_service", service)
    monkeypatch.setattr(app_module, "paper_report_service", PaperReportService(service))


def _client():
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


def test_paper_account_initializes_account(tmp_path, monkeypatch):
    service = _service(tmp_path)
    _patch_paper_services(monkeypatch, service)

    response = _client().get("/api/paper/account")

    assert response.status_code == 200
    data = response.get_json()
    assert data["initial_cash"] == 100000.0
    assert data["cash"] == 100000.0
    assert data["total_assets"] == 100000.0
    assert data["created_at"]


def test_paper_list_endpoints_start_empty(tmp_path, monkeypatch):
    service = _service(tmp_path)
    _patch_paper_services(monkeypatch, service)
    client = _client()

    assert client.get("/api/paper/positions").get_json() == []
    assert client.get("/api/paper/orders").get_json() == []
    assert client.get("/api/paper/equity").get_json() == []


def test_paper_run_executes_trade(tmp_path, monkeypatch):
    service = _service(tmp_path, quotes={"600519": _quote("600519", 100.0)})
    _patch_paper_services(monkeypatch, service)

    response = _client().post(
        "/api/paper/run",
        json={"scores": {"600519": 80}, "trading_date": "2026-06-25"},
    )

    assert response.status_code == 200
    data = response.get_json()
    assert data["results"][0]["status"] == "filled"
    assert data["results"][0]["action"] == "BUY"
    assert data["account"]["cash"] < 100000.0


def test_paper_run_duplicate_is_business_result(tmp_path, monkeypatch):
    service = _service(tmp_path, quotes={"600519": _quote("600519", 100.0)})
    service.run_scores({"600519": 80}, trading_date="2026-06-25")
    _patch_paper_services(monkeypatch, service)

    response = _client().post(
        "/api/paper/run",
        json={"scores": {"600519": 80}, "trading_date": "2026-06-25"},
    )

    assert response.status_code == 200
    assert response.get_json()["results"][0]["reason"] == "duplicate_run"


def test_paper_run_dry_run_has_no_database_side_effects(tmp_path, monkeypatch):
    service = _service(tmp_path, quotes={"600519": _quote("600519", 100.0)})
    _patch_paper_services(monkeypatch, service)

    response = _client().post(
        "/api/paper/run",
        json={"scores": {"600519": 80}, "trading_date": "2026-06-25", "dry_run": True},
    )

    assert response.status_code == 200
    data = response.get_json()
    assert data["account"] is None
    assert data["results"][0]["status"] == "dry_run"
    assert not service.db_path.exists()


def test_paper_run_rejects_missing_scores(tmp_path, monkeypatch):
    _patch_paper_services(monkeypatch, _service(tmp_path))

    response = _client().post("/api/paper/run", json={})

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_paper_run_rejects_short_stock_code(tmp_path, monkeypatch):
    _patch_paper_services(monkeypatch, _service(tmp_path))

    response = _client().post("/api/paper/run", json={"scores": {"1": 80}})

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_paper_run_rejects_invalid_score(tmp_path, monkeypatch):
    _patch_paper_services(monkeypatch, _service(tmp_path))

    response = _client().post("/api/paper/run", json={"scores": {"600519": 101}})

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_paper_run_rejects_later_invalid_score_without_partial_trade(tmp_path, monkeypatch):
    service = _service(
        tmp_path,
        quotes={
            "600519": _quote("600519", 100.0),
            "000001": _quote("000001", 10.0),
        },
    )
    _patch_paper_services(monkeypatch, service)

    response = _client().post(
        "/api/paper/run",
        json={"scores": {"600519": 80, "000001": 101}, "trading_date": "2026-06-25"},
    )

    assert response.status_code == 400
    assert "error" in response.get_json()
    assert not service.db_path.exists()


def test_paper_run_rejects_malformed_score_type(tmp_path, monkeypatch):
    _patch_paper_services(monkeypatch, _service(tmp_path))

    response = _client().post("/api/paper/run", json={"scores": {"600519": True}})

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_paper_run_rejects_invalid_date(tmp_path, monkeypatch):
    _patch_paper_services(monkeypatch, _service(tmp_path))

    response = _client().post("/api/paper/run", json={"scores": {"600519": 80}, "trading_date": "bad-date"})

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_paper_run_rejects_non_boolean_dry_run(tmp_path, monkeypatch):
    _patch_paper_services(monkeypatch, _service(tmp_path))

    response = _client().post("/api/paper/run", json={"scores": {"600519": 80}, "dry_run": "false"})

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_paper_run_quote_unavailable_returns_skipped(tmp_path, monkeypatch):
    service = _service(tmp_path, quotes={"600519": None})
    _patch_paper_services(monkeypatch, service)

    response = _client().post("/api/paper/run", json={"scores": {"600519": 80}, "trading_date": "2026-06-25"})

    assert response.status_code == 200
    result = response.get_json()["results"][0]
    assert result["status"] == "skipped"
    assert result["reason"] == "quote_unavailable"


def test_paper_orders_limit_query(tmp_path, monkeypatch):
    service = _service(tmp_path, quotes={"600519": _quote("600519", 100.0)})
    for day, score in (("2026-06-25", 80), ("2026-06-26", 50), ("2026-06-27", 20)):
        service.run_scores({"600519": score}, trading_date=day)
    _patch_paper_services(monkeypatch, service)

    response = _client().get("/api/paper/orders?limit=2")

    assert response.status_code == 200
    assert len(response.get_json()) == 2


def test_paper_report_endpoint_returns_valid_workbook(tmp_path, monkeypatch):
    service = _service(tmp_path, quotes={"600519": _quote("600519", 100.0)})
    service.run_scores({"600519": 80}, trading_date="2026-06-25")
    _patch_paper_services(monkeypatch, service)

    response = _client().get("/api/paper/report.xlsx")


    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Account", "Positions", "Orders", "Equity", "TaskRuns", "Metadata"]


def test_health_does_not_create_paper_database(tmp_path, monkeypatch):
    service = PaperTradingService(base_dir=tmp_path, db_path=tmp_path / "paper.sqlite3", config=_config())
    _patch_paper_services(monkeypatch, service)

    response = _client().get("/health")

    assert response.status_code == 200
    assert not service.db_path.exists()
