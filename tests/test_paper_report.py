"""Paper trading Excel report tests."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from core.paper_report import DISCLAIMER, PaperReportService
from tests.test_paper_trading import _quote, _service


EXPECTED_SHEETS = ["Account", "Positions", "Orders", "Equity", "TaskRuns", "Metadata"]


def _workbook(data: BytesIO):
    return load_workbook(data, read_only=True, data_only=True)


def _metadata_values(workbook) -> dict[str, str]:
    rows = workbook["Metadata"].iter_rows(min_row=2, values_only=True)
    return {str(key): "" if value is None else str(value) for key, value in rows}


def test_empty_report_workbook_has_expected_sheets(tmp_path):
    service = _service(tmp_path)
    report = PaperReportService(service)

    workbook = _workbook(report.generate_xlsx())

    assert workbook.sheetnames == EXPECTED_SHEETS
    metadata = _metadata_values(workbook)
    assert metadata["disclaimer"] == DISCLAIMER
    assert metadata["simulation_start_at"]


def test_report_workbook_contains_trade_and_metadata(tmp_path):
    service = _service(tmp_path, quotes={"600519": _quote("600519", 100.0)})
    service.run_scores({"600519": 80}, trading_date="2026-06-25")
    report = PaperReportService(service)

    workbook = _workbook(report.generate_xlsx())

    order_rows = list(workbook["Orders"].iter_rows(values_only=True))
    assert order_rows[0][3] == "stock_code"
    assert order_rows[1][3] == "600519"
    metadata = _metadata_values(workbook)
    assert metadata["disclaimer"] == DISCLAIMER
    assert metadata["data_range"] == "2026-06-25 ~ 2026-06-25"
